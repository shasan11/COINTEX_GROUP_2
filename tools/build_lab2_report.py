"""Build the COMP6009 Lab 2 report as a polished Word document."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "lab2"
REPORT_PATH = OUT / "COMP6009_Lab2_Report.docx"
CHART_PATH = OUT / "success_rate_by_level.png"

NAVY = "17365D"
BLUE = "2F75B5"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F3F5F7"
BORDER = "D9D9D9"
BLACK = RGBColor(0, 0, 0)
FONT = "Arial"


def read_csv(name: str) -> list[dict[str, str]]:
    with (OUT / name).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def set_cell_fill(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, size=10.5, bold=False, color=BLACK, italic=False) -> None:
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def add_text(doc, text: str, *, bold_lead: str | None = None, style=None, align=None, before=0, after=6):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.12
    if align is not None:
        paragraph.alignment = align
    if bold_lead and text.startswith(bold_lead):
        first = paragraph.add_run(bold_lead)
        set_run_font(first, bold=True)
        remainder = paragraph.add_run(text[len(bold_lead):])
        set_run_font(remainder)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullet(doc, text: str, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.line_spacing = 1.08
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_number(doc, text: str):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text: str, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(5)
    for run in paragraph.runs:
        set_run_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)
    return paragraph


def add_table(doc, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_fill(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(header)
        set_run_font(run, size=9.2, bold=True, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            if widths:
                cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_fill(cell, PALE_BLUE if row_index % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(str(value)), size=9.0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_caption(doc, text: str):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, italic=True)


def configure_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = BLACK
    for style_name, size in (("Title", 25), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run("COMP6009 Lab 2    |    CoinTex Genetic Algorithm")
    set_run_font(run, size=8, color=RGBColor(90, 90, 90))
    return doc


def build_report() -> Path:
    summary = read_csv("benchmark_summary.csv")
    tuning = read_csv("tuning_summary.csv")
    by_key = {(row["configuration"], int(row["level"])): row for row in summary}

    doc = configure_document()

    # Cover
    doc.add_paragraph().paragraph_format.space_after = Pt(22)
    if (ROOT / "cointex_logo.png").exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(ROOT / "cointex_logo.png"), width=Inches(1.55))
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(15)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("CoinTex Genetic Algorithm Improvement")
    set_run_font(run, size=25, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(30)
    set_run_font(subtitle.add_run("COMP6009 Lab 2 Report"), size=16, bold=True)
    for line in (
        "Student name and ID: ______________________________",
        "Group members: ____________________________________",
        "Submission date: ___________________________________",
    ):
        add_text(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_text(
        doc,
        "This report documents the submitted algorithm, controlled experiments, measured results, limitations, and proposed next improvement.",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        before=24,
        after=0,
    )
    doc.add_page_break()

    add_heading(doc, "Executive Summary", 1)
    add_text(
        doc,
        "The submitted agent replaces the Blackboard program's unsafe search ranges, ineffective hazard penalties, and ten-member random-selection population with a bounded real-valued genetic algorithm. The improved fitness evaluates every remaining coin, measures danger along the full movement path, increases progress pressure when time is short, and rewards escape from nearby hazards. Tournament selection, elitism, blend-style crossover, and adaptive mutation maintain useful solutions while preserving exploration.",
    )
    add_text(
        doc,
        "Across 320 deterministic trials, the baseline achieved 25.0% success and the improved configuration achieved 71.25%. On level 40, success increased from 0% to 85%. The improvement did not solve the hardest simulated level reliably: level 60 success was 10%, and mean remaining health was below 1%. This result supports a hybrid next step that combines the GA's global target search with short-horizon collision prediction.",
    )
    add_heading(doc, "Deliverables", 2)
    add_bullet(doc, "Improved Blackboard-compatible program in main_updated.py")
    add_bullet(doc, "Reusable GA implementation and fixed-seed benchmark in PlayerGA")
    add_bullet(doc, "Raw trial data, tuning results, summary tables, and KPI chart")
    add_bullet(doc, "Presentation deck structured for a 5 to 7 minute group presentation")

    add_heading(doc, "Part 1 Understanding the Program and Genetic Algorithm Design", 1)
    add_heading(doc, "Program Structure", 2)
    add_text(
        doc,
        "The tutorial application uses Kivy for the interface and PyGAD for optimization. Each level owns widgets for the player, coins, monsters, and fires. When a level starts, CollectCoinThread creates a PyGAD instance in a background thread. The fitness callback reads the current widget coordinates, and the generation callback sends the best chromosome to start_char_animation. Kivy then moves the character and checks collisions and coin collection.",
    )
    add_table(
        doc,
        ["Component", "Responsibility"],
        [
            ["main_updated.py", "Connects the tutorial GUI to PyGAD and launches the 24-level application"],
            ["PlayerGA/ga_core.py", "Contains testable baseline and improved fitness plus evolutionary operators"],
            ["PlayerGA/benchmark.py", "Runs controlled trials and records success, time, health, and coin completion"],
            ["cointex.kv", "Defines the original tutorial screens, controls, and level widgets"],
        ],
        [1.65, 5.15],
    )

    add_heading(doc, "Chromosome Representation", 2)
    add_text(
        doc,
        "A chromosome contains two floating-point genes, [x, y]. Both values describe a target in normalized screen coordinates. The controller converts the best chromosome into a Kivy movement target. The chromosome does not encode a complete route. The GA therefore solves a repeated local planning problem: evolve a target, move toward it, observe the changed game state, and search again.",
    )
    add_table(
        doc,
        ["Step", "Operation", "Game effect"],
        [
            ["1", "Sense", "Read player, coin, monster, and fire coordinates"],
            ["2", "Evaluate", "Score each [x, y] target using progress and path safety"],
            ["3", "Select and reproduce", "Keep strong targets and create varied offspring"],
            ["4", "Act", "Move toward the best bounded target"],
            ["5", "Repeat", "Re-evaluate after the player and hazards move"],
        ],
        [0.55, 1.85, 4.40],
    )
    add_caption(doc, "Figure 1  Receding-target genetic control workflow")

    add_heading(doc, "Original Fitness and GA Settings", 2)
    add_text(
        doc,
        "The supplied Blackboard file measured Manhattan distance to the nearest coin, but its monster and fire branches added or subtracted zero. It therefore rewarded coin proximity without effective safety control. The initial range extended to 1.3 even though visible coordinates end at 1.0, and replacement mutations used values from 1.0 to 10.0. Swap mutation also has little meaning for a two-gene position because it only exchanges x and y. A ten-member population with random parent selection provided little diversity or selection pressure.",
    )
    add_table(
        doc,
        ["Design choice", "Blackboard version", "Submitted version"],
        [
            ["Population", "10", "60"],
            ["Parents", "2, random selection", "15, tournament selection of 4"],
            ["Coordinate range", "Initial 0.3 to 1.3; mutation 1 to 10", "Playable x 0.07 to 0.93; y 0.12 to 0.90"],
            ["Crossover", "PyGAD default", "Scattered in GUI; blend in benchmark core"],
            ["Mutation", "Swap, 10% of genes", "Adaptive, 0.60 for weak and 0.20 for strong chromosomes"],
            ["Elitism", "None specified", "Best 3 chromosomes retained"],
            ["Hazard fitness", "Zero effective penalty", "Path-clearance penalty with danger radius"],
        ],
        [1.42, 2.52, 2.86],
    )

    add_heading(doc, "Part 2 Algorithm Improvement and Parameter Tuning", 1)
    add_heading(doc, "Improved Fitness Function", 2)
    add_text(doc, "The submitted fitness is the sum of four interpretable terms and one boundary rule:")
    add_bullet(doc, "Coin progress considers every coin and favors a reachable nearby goal rather than dictionary order.")
    add_bullet(doc, "Path safety measures the shortest gap between each danger and the full segment from player to target.")
    add_bullet(doc, "Urgency increases coin attraction as the level timer falls, while retaining a safety penalty.")
    add_bullet(doc, "Escape reward favors targets farther from a danger already within the player's flee radius.")
    add_bullet(doc, "Movement cost mildly favors shorter targets so the controller can react sooner.")
    add_text(
        doc,
        "Targets outside the playable area receive a large negative score. This prevents the GA from selecting coordinates that Kivy cannot use safely. The same pure function serves the tutorial launcher, the current 60-level player, the unit tests, and the benchmark.",
    )

    add_heading(doc, "Experiment Design", 2)
    add_text(
        doc,
        "The controlled benchmark used current CoinTex level parameters in a deterministic two-dimensional simulator. Levels 1, 20, 40, and 60 represent increasing coin count, enemy speed, chase behavior, contact damage, and time pressure. Each configuration ran with 20 fixed seeds per level. Four configurations separated the contribution of the fitness function from the GA operators. Both main comparison configurations used the same 60-member population, so the success comparison does not give the improved agent a larger evaluation budget.",
    )
    add_table(
        doc,
        ["Configuration", "Fitness", "Operators"],
        [
            ["Baseline", "Original first-coin and square danger rule", "Uniform crossover, 0.10 mutation, no elitism"],
            ["Fitness only", "Improved multi-coin path-safety fitness", "Baseline operators"],
            ["Operators only", "Baseline fitness", "Tournament, elitism, blend, adaptive mutation"],
            ["Improved", "Improved multi-coin path-safety fitness", "Improved operators"],
        ],
        [1.30, 2.85, 2.65],
    )
    add_text(
        doc,
        "The benchmark recorded goal success, elapsed simulated time, remaining health percentage, and coin completion. It produced 320 comparison trials. A separate grid tested populations of 60, 120, and 240 with mutation probabilities of 0.15, 0.30, and 0.50. Each tuning candidate ran 16 trials across levels 20 and 40.",
    )

    add_heading(doc, "Main Results", 2)
    result_rows = []
    for level_number in (1, 20, 40, 60):
        base = by_key[("baseline", level_number)]
        improved = by_key[("improved", level_number)]
        result_rows.append([
            str(level_number),
            base["success_rate_pct"] + "%",
            improved["success_rate_pct"] + "%",
            improved["mean_success_time_seconds"] or "No wins",
            improved["mean_health_pct"] + "%",
            improved["mean_coin_completion_pct"] + "%",
        ])
    add_table(
        doc,
        ["Level", "Baseline success", "Improved success", "Improved win time s", "Improved health", "Improved coins"],
        result_rows,
        [0.55, 1.12, 1.12, 1.25, 1.05, 1.05],
    )
    if CHART_PATH.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(CHART_PATH), width=Inches(6.65))
        add_caption(doc, "Figure 2  Success rate from 20 fixed-seed trials per configuration and level")
    add_text(
        doc,
        "Averaged across the four tested levels, baseline success was 25.0% and improved success was 71.25%, a gain of 46.25 percentage points. The fitness-only configuration reached 57.5%, while operators alone reached 36.25%. The ablation therefore attributes most of the gain to the new fitness function. The improved operators added speed and stability: at level 20, the full system won in 13.2 seconds on average, compared with 21.39 seconds for fitness-only wins.",
    )

    add_heading(doc, "Parameter Tuning", 2)
    tuning_rows = []
    for row in tuning:
        tuning_rows.append([
            row["population_size"],
            row["mutation_probability"],
            row["success_rate_pct"] + "%",
            (row["mean_success_time_seconds"] or "-") + (" s" if row["mean_success_time_seconds"] else ""),
            row["mean_health_pct"] + "%",
        ])
    add_table(
        doc,
        ["Population", "Mutation", "Success", "Win time", "Health"],
        tuning_rows,
        [1.15, 1.15, 1.15, 1.50, 1.20],
    )
    add_text(
        doc,
        "The 60-member population with 0.30 mutation was the only tested candidate to win all 16 tuning trials. It also retained 57.86% mean health. Larger populations sometimes reduced win time, but they required two to four times as many fitness evaluations and did not improve success in this sample. The submitted GUI therefore uses population 60, 15 parents, 3 elites, and adaptive mutation centered on the 0.30 exploration rate.",
    )

    add_heading(doc, "Part 3 Performance Analysis and Further Improvements", 1)
    add_heading(doc, "Why the Changes Worked", 2)
    add_text(
        doc,
        "The original objective could select a coin through a monster or fire because it scored only the final point and the hazard branches had zero effect. The new segment-clearance term changes that geometry. A target behind a hazard receives a penalty even when the endpoint itself is safe. Evaluating every coin also removes sensitivity to dictionary order. These two changes explain why the fitness-only ablation produced the largest improvement.",
    )
    add_text(
        doc,
        "Tournament selection and elitism prevent useful safe targets from disappearing after one generation. Blend crossover creates intermediate coordinates instead of merely swapping the two parent genes. Adaptive mutation restores exploration when the population stops improving. These operator changes reduce the time spent oscillating between similar targets, which is visible in the shorter successful-run times for the full system.",
    )

    add_heading(doc, "Failure on the Hardest Level", 2)
    add_text(
        doc,
        "Level 60 remained unsolved in 90% of simulated trials. The improved agent still collected 69.0% of coins on average, but mean remaining health was 0.97%. The two-gene chromosome explains this limitation. It represents only the next endpoint and cannot encode a multi-step route around moving dangers. The straight-line path estimate also assumes a danger's near-term motion from a single snapshot. Fast chasers and pulsing fire can invalidate that estimate before the character reaches the target.",
    )
    add_text(
        doc,
        "The benchmark is a controlled simulator based on the current level parameters. It supports reproducible comparison, but it does not prove identical success in the event-driven Kivy GUI. Final submission testing should include recorded GUI attempts on the selected tutorial levels, using the same success, time, and remaining-life measures.",
    )

    add_heading(doc, "Proposed Hybrid Improvement", 2)
    add_text(
        doc,
        "A practical next version would keep the GA for global target selection and add a short-horizon model-predictive safety layer. After the GA proposes [x, y], the safety layer would simulate several small movement steps, predict danger positions, and reject a target if any predicted clearance falls below the collision radius. It would then choose the safest nearby waypoint. This hybrid can represent a curved route as a sequence of accepted waypoints without expanding the chromosome into a long path. The method should help on level 60 because it reacts to motion during travel, which is the main information missing from the present fitness.",
    )
    add_text(
        doc,
        "Reinforcement learning is another option, but it would require many training episodes, a carefully designed reward, and a stable simulator. The hybrid planner is the stronger immediate proposal because it reuses the tested GA and directly addresses the observed collision failure.",
    )

    add_heading(doc, "Reproduction and Code Submission", 1)
    add_text(doc, "From the repository root, install the research dependencies and run:")
    for command in (
        "python main_updated.py",
        "python PlayerGA/benchmark.py --output output/lab2 --seeds 20 --levels 1 20 40 60",
        "python -m unittest discover -s tests -v",
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(4)
        run = paragraph.add_run(command)
        set_run_font(run, size=9.5)
        run.font.name = "Courier New"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Courier New")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Courier New")
    add_text(
        doc,
        "The output/lab2 folder contains benchmark_trials.csv, benchmark_summary.csv, tuning_trials.csv, tuning_summary.csv, benchmark_metadata.json, and the success-rate chart. Fixed seeds make the trial layouts and GA populations repeatable.",
    )

    add_heading(doc, "References", 1)
    references = [
        "COMP6009 Lab Task 2, version 3, last modified 4 September 2026.",
        "Gad, A. F. (2021). PyGAD: An Intuitive Genetic Algorithm Python Library. arXiv:2106.06158.",
        "Gad, A. F. CoinTex source repository. https://github.com/ahmedfgad/CoinTex",
        "PyGAD documentation. https://pygad.readthedocs.io/",
    ]
    for item in references:
        add_text(doc, item, after=4)

    doc.core_properties.title = "CoinTex Genetic Algorithm Improvement"
    doc.core_properties.subject = "COMP6009 Lab 2"
    doc.core_properties.author = "Student"
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(REPORT_PATH)
    return REPORT_PATH


if __name__ == "__main__":
    print(build_report())
